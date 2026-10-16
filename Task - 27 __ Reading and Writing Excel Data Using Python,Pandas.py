"""
Task 27: Reading and Writing Excel Data
Reads an Excel workbook with Pandas, inspects structure, cleans data,
creates a department summary, and exports processed results to Excel and CSV.
"""
from pathlib import Path
import pandas as pd

INPUT_FILE = Path("employee_dataset.xlsx")
OUTPUT_FILE = Path("processed_employee_data.xlsx")
OUTPUT_CSV = Path("processed_employees.csv")

def load_data(file_path: Path) -> pd.DataFrame:
    """Read the Employees sheet from an Excel workbook."""
    return pd.read_excel(file_path, sheet_name="Employees", engine="openpyxl")

def clean_data(df: pd.DataFrame) -> pd.DataFrame:
    """Clean common data-quality issues and return a new DataFrame."""
    data = df.copy()

    # Standardize column labels and text fields
    data.columns = data.columns.str.strip().str.replace(" ", "_")
    for column in ["Employee_Name", "Department", "Location", "Email"]:
        if column in data.columns:
            data[column] = data[column].astype("string").str.strip()

    if "Employee_Name" in data.columns:
        data["Employee_Name"] = data["Employee_Name"].str.title()
    if "Department" in data.columns:
        data["Department"] = data["Department"].replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
        data["Department"] = data["Department"].fillna("Unknown")
        data["Department"] = data["Department"].str.title()
    if "Location" in data.columns:
        data["Location"] = data["Location"].str.title()

    # Convert currency-like salary values to numeric
    if "Salary" in data.columns:
        salary_text = data["Salary"].astype("string").str.replace(r"[₹,$,]", "", regex=True)
        data["Salary"] = pd.to_numeric(salary_text, errors="coerce")

    # Convert dates and ages to usable types
    if "Joining_Date" in data.columns:
        data["Joining_Date"] = pd.to_datetime(data["Joining_Date"], errors="coerce")
    if "Age" in data.columns:
        data["Age"] = pd.to_numeric(data["Age"], errors="coerce")

    # Remove exact duplicate records
    data = data.drop_duplicates().reset_index(drop=True)

    # Fill missing age with the median; keep missing email as a blank value
    if "Age" in data.columns and data["Age"].isna().any():
        data["Age"] = data["Age"].fillna(data["Age"].median()).round().astype("Int64")
    if "Email" in data.columns:
        data["Email"] = data["Email"].fillna("")

    return data

def make_department_summary(data: pd.DataFrame) -> pd.DataFrame:
    """Aggregate employee count and salary statistics by department."""
    summary = (
        data.groupby("Department", dropna=False)
        .agg(
            Employee_Count=("Employee_ID", "count"),
            Average_Salary=("Salary", "mean"),
            Median_Salary=("Salary", "median"),
            Minimum_Salary=("Salary", "min"),
            Maximum_Salary=("Salary", "max"),
        )
        .reset_index()
        .sort_values("Employee_Count", ascending=False)
    )
    salary_columns = ["Average_Salary", "Median_Salary", "Minimum_Salary", "Maximum_Salary"]
    summary[salary_columns] = summary[salary_columns].round(2)
    return summary

def export_results(data: pd.DataFrame, summary: pd.DataFrame) -> None:
    """Write cleaned data and summary sheets to an Excel workbook."""
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        data.to_excel(writer, sheet_name="Cleaned_Employees", index=False)
        summary.to_excel(writer, sheet_name="Department_Summary", index=False)
        data.describe(include="all").transpose().to_excel(writer, sheet_name="Data_Profile")
    data.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

    # Add readable formatting and filters to exported workbook
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = load_workbook(OUTPUT_FILE)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203864")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for col in worksheet.columns:
            max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col)
            worksheet.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 32)
    workbook.save(OUTPUT_FILE)

def main() -> None:
    raw = load_data(INPUT_FILE)
    print("RAW SHAPE:", raw.shape)
    print("COLUMNS:", raw.columns.tolist())
    print("MISSING VALUES BEFORE CLEANING:\\n", raw.isna().sum())
    cleaned = clean_data(raw)
    summary = make_department_summary(cleaned)
    print("\\nCLEANED SHAPE:", cleaned.shape)
    print("\\nDEPARTMENT SUMMARY:\\n", summary.to_string(index=False))
    export_results(cleaned, summary)
    print(f"\\nSaved: {OUTPUT_FILE.resolve()}")
    print(f"Saved: {OUTPUT_CSV.resolve()}")

if __name__ == "__main__":
    main()
